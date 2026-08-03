"""Excel export (OpenXML via openpyxl).

Produces one workbook with three sheets:
  Orders     — one row per order line, the format the distributor asked for
  Load Sheet — product-wise box totals for the warehouse purchase order
  Summary    — day/range headline numbers

Money is written as real numbers with an Indian rupee number format, not as
pre-formatted strings, so the admin can pivot and re-total in Excel.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import models
from . import reports

RUPEE_FMT = '₹#,##0.00'
INT_FMT = '#,##0'

HEADER_FILL = PatternFill("solid", fgColor="1E7D3A")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TOTAL_FONT = Font(bold=True)
THIN = Side(style="thin", color="D8DEE4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _rupees(paise: int) -> float:
    return round(paise / 100, 2)


def _write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.freeze_panes = "A2"


def _autosize(ws, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def build_workbook(db: Session, from_date: date, to_date: date) -> bytes:
    wb = Workbook()

    # ---------------- Orders ----------------
    ws = wb.active
    ws.title = "Orders"
    _write_header(
        ws,
        [
            "Date",
            "Customer Name",
            "Mobile Number",
            "Product Name",
            "MRP",
            "Box Price",
            "Quantity (Boxes)",
            "Line Total",
            "Grand Total",
            "Salesman Name",
        ],
    )

    orders = list(
        db.scalars(
            select(models.Order)
            .where(
                models.Order.is_deleted.is_(False),
                models.Order.order_date >= from_date,
                models.Order.order_date <= to_date,
            )
            .order_by(models.Order.order_date, models.Order.booked_at, models.Order.id)
        ).unique()
    )

    grand_boxes = 0
    grand_value = 0
    for order in orders:
        salesman = order.salesman.name if order.salesman else ""
        for position, item in enumerate(order.items):
            ws.append(
                [
                    order.order_date,
                    order.customer.name,
                    order.customer.mobile,
                    item.product_name,
                    _rupees(item.mrp_paise),
                    _rupees(item.box_price_paise),
                    item.qty_boxes,
                    _rupees(item.line_total_paise),
                    # Repeating the order total on every line would break SUM();
                    # it belongs on the order's first line only.
                    _rupees(order.total_value_paise) if position == 0 else None,
                    salesman,
                ]
            )
        grand_boxes += order.total_boxes
        grand_value += order.total_value_paise

    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "dd-mmm-yyyy"
        for col in (4, 5, 7, 8):
            row[col].number_format = RUPEE_FMT
        row[6].number_format = INT_FMT

    if orders:
        ws.append([])
        total_row = ws.max_row + 1
        ws.cell(row=total_row, column=3, value="TOTAL").font = TOTAL_FONT
        boxes_cell = ws.cell(row=total_row, column=7, value=grand_boxes)
        value_cell = ws.cell(row=total_row, column=9, value=_rupees(grand_value))
        boxes_cell.font = TOTAL_FONT
        boxes_cell.number_format = INT_FMT
        value_cell.font = TOTAL_FONT
        value_cell.number_format = RUPEE_FMT

    _autosize(ws, [12, 26, 15, 30, 10, 12, 16, 14, 14, 18])
    ws.auto_filter.ref = f"A1:J{max(ws.max_row, 1)}"

    # ---------------- Load Sheet ----------------
    sheet = reports.load_sheet(db, from_date, to_date)
    ws2 = wb.create_sheet("Load Sheet")
    _write_header(ws2, ["Product Name", "Total Boxes", "Total Value"])
    for line in sheet.lines:
        ws2.append([line.product_name, line.total_boxes, _rupees(line.total_value_paise)])
    for row in ws2.iter_rows(min_row=2):
        row[1].number_format = INT_FMT
        row[2].number_format = RUPEE_FMT
    if sheet.lines:
        ws2.append([])
        total_row = ws2.max_row + 1
        ws2.cell(row=total_row, column=1, value="TOTAL").font = TOTAL_FONT
        c = ws2.cell(row=total_row, column=2, value=sheet.total_boxes)
        c.font = TOTAL_FONT
        c.number_format = INT_FMT
        c = ws2.cell(
            row=total_row,
            column=3,
            value=_rupees(sum(line.total_value_paise for line in sheet.lines)),
        )
        c.font = TOTAL_FONT
        c.number_format = RUPEE_FMT
    _autosize(ws2, [34, 14, 16])

    # ---------------- Summary ----------------
    data = reports.summary(db, from_date, to_date)
    ws3 = wb.create_sheet("Summary")
    _write_header(ws3, ["Metric", "Value"])
    ws3.append(["From", from_date])
    ws3.append(["To", to_date])
    ws3.append(["Total Orders", data.total_orders])
    ws3.append(["Total Customers", data.total_customers])
    ws3.append(["Total Boxes", data.total_boxes])
    ws3.append(["Total Value", _rupees(data.total_value_paise)])
    ws3["B2"].number_format = "dd-mmm-yyyy"
    ws3["B3"].number_format = "dd-mmm-yyyy"
    ws3["B7"].number_format = RUPEE_FMT

    ws3.append([])
    ws3.append(["Salesman", "Orders", "Boxes", "Value"])
    for cell in ws3[ws3.max_row]:
        cell.font = TOTAL_FONT
    for line in data.salesman_wise:
        ws3.append(
            [line.salesman_name, line.orders, line.total_boxes, _rupees(line.total_value_paise)]
        )
        ws3.cell(row=ws3.max_row, column=4).number_format = RUPEE_FMT
    _autosize(ws3, [24, 16, 14, 16])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def filename_for(from_date: date, to_date: date) -> str:
    if from_date == to_date:
        return f"ecoo-orders-{from_date.isoformat()}.xlsx"
    return f"ecoo-orders-{from_date.isoformat()}-to-{to_date.isoformat()}.xlsx"
