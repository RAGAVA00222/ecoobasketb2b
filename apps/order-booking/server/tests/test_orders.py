from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook


def book(client, catalogue, *, uuid, mobile, name, qty_a=5, qty_b=0, salesman="SM01"):
    a, b = catalogue
    items = [{"product_id": a["id"], "qty_boxes": qty_a}]
    if qty_b:
        items.append({"product_id": b["id"], "qty_boxes": qty_b})
    return client.post(
        "/api/v1/orders",
        json={
            "client_uuid": uuid,
            "mobile": mobile,
            "customer_name": name,
            "salesman_code": salesman,
            "items": items,
        },
    )


def test_order_totals_are_computed_server_side(client, catalogue):
    resp = book(client, catalogue, uuid="uuid-order-0001", mobile="9876543210", name="Sri Stores")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["total_boxes"] == 5
    assert body["total_products"] == 1
    assert body["total_value_paise"] == 5 * 92000  # ₹4,600.00
    assert body["items"][0]["line_total_paise"] == 460000


def test_mobile_is_normalised_so_one_shop_is_one_customer(client, catalogue):
    first = book(client, catalogue, uuid="uuid-order-0001", mobile="+91 98765 43210", name="Sri Stores")
    second = book(client, catalogue, uuid="uuid-order-0002", mobile="09876543210", name="Sri Stores")
    assert first.json()["customer_id"] == second.json()["customer_id"]
    assert first.json()["customer_mobile"] == "9876543210"


def test_customer_lookup_autofills_name(client, catalogue):
    book(client, catalogue, uuid="uuid-order-0001", mobile="9876543210", name="Sri Stores")
    resp = client.get("/api/v1/customers/lookup", params={"mobile": "9876543210"})
    assert resp.status_code == 200
    assert resp.json()["name"] == "Sri Stores"
    assert client.get(
        "/api/v1/customers/lookup", params={"mobile": "9000000000"}
    ).status_code == 404


def test_resync_of_the_same_order_does_not_double_book(client, catalogue):
    a, _ = catalogue
    payload = {
        "client_uuid": "uuid-order-dup1",
        "mobile": "9876543210",
        "customer_name": "Sri Stores",
        "items": [{"product_id": a["id"], "qty_boxes": 3}],
    }
    first = client.post("/api/v1/orders/sync", json=[payload]).json()
    second = client.post("/api/v1/orders/sync", json=[payload]).json()
    assert first["results"][0]["status"] == "created"
    assert second["results"][0]["status"] == "duplicate"
    assert first["results"][0]["order_id"] == second["results"][0]["order_id"]
    assert len(client.get("/api/v1/orders").json()) == 1


def test_one_bad_row_does_not_block_the_rest_of_the_batch(client, catalogue):
    a, _ = catalogue
    good = {
        "client_uuid": "uuid-order-good",
        "mobile": "9876543210",
        "customer_name": "Sri Stores",
        "items": [{"product_id": a["id"], "qty_boxes": 2}],
    }
    bad = {
        "client_uuid": "uuid-order-bad0",
        "mobile": "9876543211",
        "customer_name": "Ghost Stores",
        "items": [{"product_id": 999999, "qty_boxes": 2}],
    }
    results = client.post("/api/v1/orders/sync", json=[bad, good]).json()["results"]
    by_uuid = {r["client_uuid"]: r for r in results}
    assert by_uuid["uuid-order-bad0"]["status"] == "rejected"
    assert by_uuid["uuid-order-good"]["status"] == "created"


def test_repeated_product_lines_are_merged(client, catalogue):
    a, _ = catalogue
    resp = client.post(
        "/api/v1/orders",
        json={
            "client_uuid": "uuid-order-mrg1",
            "mobile": "9876543210",
            "customer_name": "Sri Stores",
            "items": [
                {"product_id": a["id"], "qty_boxes": 2},
                {"product_id": a["id"], "qty_boxes": 3},
            ],
        },
    )
    body = resp.json()
    assert len(body["items"]) == 1
    assert body["items"][0]["qty_boxes"] == 5
    assert body["total_boxes"] == 5


def test_price_change_does_not_repricebooked_orders(client, catalogue, admin_headers):
    a, _ = catalogue
    booked = book(client, catalogue, uuid="uuid-order-0001", mobile="9876543210", name="Sri Stores")
    client.patch(
        f"/api/v1/products/{a['id']}", headers=admin_headers, json={"box_price_paise": 99000}
    )
    after = client.get(f"/api/v1/orders/{booked.json()['id']}").json()
    assert after["items"][0]["box_price_paise"] == 92000
    assert after["total_value_paise"] == 5 * 92000


def test_deleted_order_leaves_the_load_sheet(client, catalogue):
    order = book(client, catalogue, uuid="uuid-order-0001", mobile="9876543210", name="Sri Stores")
    assert client.get("/api/v1/reports/load-sheet").json()["total_boxes"] == 5
    client.delete(f"/api/v1/orders/{order.json()['id']}")
    assert client.get("/api/v1/reports/load-sheet").json()["total_boxes"] == 0
    assert client.get("/api/v1/orders").json() == []


def test_load_sheet_aggregates_boxes_across_customers(client, catalogue):
    book(client, catalogue, uuid="uuid-order-0001", mobile="9876543210", name="Sri Stores", qty_a=5, qty_b=3)
    book(client, catalogue, uuid="uuid-order-0002", mobile="9876543211", name="Amman Stores", qty_a=7)
    sheet = client.get("/api/v1/reports/load-sheet").json()
    lines = {line["product_name"]: line["total_boxes"] for line in sheet["lines"]}
    assert lines == {"Good Day Biscuit": 12, "Milk Bikis": 3}
    assert sheet["total_boxes"] == 15
    # Ordered most boxes first — that is the order the warehouse reads it in.
    assert sheet["lines"][0]["product_name"] == "Good Day Biscuit"


def test_summary_counts_each_shop_once(client, catalogue):
    book(client, catalogue, uuid="uuid-order-0001", mobile="9876543210", name="Sri Stores")
    book(client, catalogue, uuid="uuid-order-0002", mobile="9876543210", name="Sri Stores")
    summary = client.get("/api/v1/reports/summary").json()
    assert summary["total_orders"] == 2
    assert summary["total_customers"] == 1
    assert summary["salesman_wise"][0]["salesman_code"] == "SM01"
    assert summary["salesman_wise"][0]["orders"] == 2


def test_excel_export_has_orders_load_sheet_and_summary(client, catalogue):
    book(client, catalogue, uuid="uuid-order-0001", mobile="9876543210", name="Sri Stores", qty_a=5, qty_b=3)
    resp = client.get("/api/v1/export/orders.xlsx")
    assert resp.status_code == 200
    assert "attachment" in resp.headers["content-disposition"]

    wb = load_workbook(BytesIO(resp.content))
    assert wb.sheetnames == ["Orders", "Load Sheet", "Summary"]

    orders = wb["Orders"]
    header = [c.value for c in orders[1]]
    assert header[:4] == ["Date", "Customer Name", "Mobile Number", "Product Name"]
    rows = [[c.value for c in row] for row in orders.iter_rows(min_row=2) if row[1].value]
    assert len(rows) == 2  # two product lines, one order
    assert rows[0][7] == 4600.00  # line total in rupees
    assert rows[0][8] == 7240.00  # grand total on the first line only
    assert rows[1][8] is None

    load = wb["Load Sheet"]
    assert [load.cell(row=2, column=c).value for c in (1, 2)] == ["Good Day Biscuit", 5]


def test_prices_are_admin_only(client):
    payload = {
        "sku": "X-1",
        "name": "Hacked Product",
        "mrp_paise": 100,
        "box_price_paise": 1,
    }
    assert client.post("/api/v1/products", json=payload).status_code == 401
    assert client.post(
        "/api/v1/products", headers={"X-Admin-Key": "wrong"}, json=payload
    ).status_code == 401


def test_deactivated_product_disappears_from_the_salesman_catalogue(
    client, catalogue, admin_headers
):
    a, _ = catalogue
    client.delete(f"/api/v1/products/{a['id']}", headers=admin_headers)
    skus = {p["sku"] for p in client.get("/api/v1/products").json()}
    assert skus == {"BIS-MILKB"}
    # ...but a delta sync still learns it was deactivated.
    delta = client.get(
        "/api/v1/products", params={"updated_since": "2000-01-01T00:00:00"}
    ).json()
    assert any(p["sku"] == "BIS-GDAY" and p["active"] is False for p in delta)
